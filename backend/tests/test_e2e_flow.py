"""API walk from upload through a saved version. The drawing is synthetic."""

from __future__ import annotations

import io
from pathlib import Path

import fitz
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import create_app
from tests.test_jobs import post_pair, settings_for


def _plant_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page(width=420, height=240)
    page.draw_rect(fitz.Rect(20, 30, 120, 110), width=1.5)
    page.insert_text((36, 78), "T-101", fontsize=12)
    page.draw_circle(fitz.Point(280, 70), 18, width=1.2)
    page.insert_text((266, 74), "P-201", fontsize=10)
    page.draw_line((120, 70), (262, 70), width=1.5)
    page.draw_circle(fitz.Point(190, 160), 16, width=1)
    page.insert_text((172, 164), "PT-101", fontsize=8)
    page.draw_line((190, 144), (190, 70), width=0.4, dashes="[2 2] 0")
    payload = document.tobytes()
    document.close()
    return payload


def _io() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "IO"
    sheet.append(["Tag", "Description"])
    sheet.append(["T-101", "Tank"])
    sheet.append(["P-201", "Pump"])
    sheet.append(["PT-101", "Pressure"])
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def test_upload_process_export_and_version(tmp_path: Path) -> None:
    application = create_app(settings_for(tmp_path, cache_dir=str(tmp_path / "cache")))
    with TestClient(application) as client:
        created = post_pair(client, pdf=_plant_pdf(), workbook=_io(), name="Synthetic plant")
        assert created.status_code == 201, created.text
        job_id = created.json()["id"]
        checked = client.post(f"/api/jobs/{job_id}/io-list/validate", json={})
        assert checked.status_code == 200, checked.text
        assert checked.json()["status"] in {"ok", "validated"}
        started = client.post(f"/api/jobs/{job_id}/process")
        assert started.status_code == 202, started.text
        finished = client.get(f"/api/jobs/{job_id}")
        failed = {
            stage["name"]: stage["error"]
            for stage in finished.json()["stages"]
            if stage["status"] == "failed"
        }
        assert not failed, failed
        assert all(stage["status"] != "running" for stage in finished.json()["stages"])
        tag_report = client.get(f"/api/jobs/{job_id}/tag-report")
        connections = client.get(f"/api/jobs/{job_id}/reports/connections")
        graph = client.get(f"/api/jobs/{job_id}/graph")
        assert tag_report.status_code == 200
        assert tag_report.json()["summary"]["total"] >= 1
        assert connections.status_code == 200
        assert graph.status_code == 200
        exports = {}
        for fmt, media in {
            "svg": "image/svg+xml",
            "json": "application/json",
            "png": "image/png",
            "pdf": "application/pdf",
            "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }.items():
            response = client.get(f"/api/jobs/{job_id}/exports/{fmt}")
            assert response.status_code == 200, response.text
            assert response.headers["content-type"].startswith(media)
            assert "attachment;" in response.headers["content-disposition"]
            assert len(response.content) > 20
            exports[fmt] = response.content
        workbook_file = client.get(f"/api/jobs/{job_id}/connections.xlsx")
        printed = client.get(f"/api/jobs/{job_id}/print.pdf")
        assert workbook_file.status_code == 200
        assert printed.status_code == 200
        assert printed.headers["content-type"].startswith("application/pdf")
        assert printed.content.startswith(b"%PDF")
        assert b"Legend" in printed.content
        svg = exports["svg"].decode("utf-8")
        assert "T-101" in svg
        assert "id=\"node-" in svg or "id='node-" in svg
        pdf = fitz.open(stream=exports["pdf"], filetype="pdf")
        assert pdf[0].get_text().strip()
        pdf.close()
        book = load_workbook(io.BytesIO(exports["excel"]), data_only=True)
        assert "Tags" in book.sheetnames
        assert "Connections" in book.sheetnames
        book.close()
        links = load_workbook(io.BytesIO(workbook_file.content), data_only=True)
        assert "Connections" in links.sheetnames
        links.close()
        draft = client.post(f"/api/jobs/{job_id}/draft")
        assert draft.status_code == 200, draft.text
        versions = client.get(f"/api/jobs/{job_id}/versions")
        original = next(item for item in versions.json() if item["number"] == 0)
        saved = client.post(f"/api/jobs/{job_id}/draft/save", json={"note": "e2e save"})
        assert saved.status_code == 200, saved.text
        assert saved.json()["number"] == 1
        again = client.get(f"/api/jobs/{job_id}/versions")
        kept = next(item for item in again.json() if item["number"] == 0)
        assert kept["svg"] == original["svg"]
        assert kept["number"] == 0
