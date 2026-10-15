"""Read sheet names, column names, and row counts. This is not tag matching."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class SheetFacts:
    """One worksheet as stored, including its header row in the row count."""

    name: str
    columns: list[str]
    row_count: int


def inspect_workbook(path: Path) -> list[SheetFacts]:
    """Open an xlsx file. The caller turns a failure into FILE_UNREADABLE."""
    return _sheets(load_workbook(path, read_only=True, data_only=True))


def inspect_workbook_bytes(data: bytes) -> list[SheetFacts]:
    """Open xlsx bytes. Nothing is written to disk."""
    return _sheets(load_workbook(BytesIO(data), read_only=True, data_only=True))


def _sheets(workbook) -> list[SheetFacts]:
    try:
        sheets: list[SheetFacts] = []
        for name in workbook.sheetnames:
            worksheet = workbook[name]
            columns: list[str] = []
            row_count = 0
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                row_count += 1
                if index == 0:
                    columns = ["" if cell is None else str(cell) for cell in row]
            sheets.append(SheetFacts(name=name, columns=columns, row_count=row_count))
        return sheets
    finally:
        workbook.close()
