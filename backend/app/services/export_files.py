"""Build downloadable files from stored results. Original uploads are not changed.

PNG and PDF are rendered with PyMuPDF from the SVG. CairoSVG is not required.
A PDF page is the SVG drawing plus the same disclaimer text.
"""

from __future__ import annotations

import hashlib
import json
import logging
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import fitz
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app import graph_config as cfg
from app.errors import AppError
from app.schemas.export import ExportMetadata
from app.schemas.graph import DiagramGraph
from app.services.storage import PDF_NAME, job_folder

logger = logging.getLogger("untangle.export")

MAX_EXPORT_BYTES = 20_000_000
CERTIFICATION_WARNING = "This file is not an engineering or safety certification."


def excel_cell(value: object) -> object:
    """Stop a cell from being read as a formula."""
    if isinstance(value, str) and value[:1] in "=+-@":
        return "'" + value
    return value


def export_metadata(
    job_id: str,
    pdf_name: str,
    pdf_hash: str,
    io_name: str,
    io_hash: str,
    version: str,
    draft: bool,
    verification_state: str,
    failed_check: bool,
    hidden_count: int,
    export_format: str,
    filename: str,
    media_type: str,
    byte_size: int,
) -> ExportMetadata:
    return ExportMetadata(
        id=hashlib.sha256(f"{job_id}|{filename}".encode()).hexdigest()[:16],
        job_id=job_id,
        version_number=0,
        format=export_format,  # type: ignore[arg-type]
        filename=filename,
        created_at=datetime.now(timezone.utc).replace(microsecond=0),
        media_type=media_type,
        byte_size=byte_size,
        source_pdf_name=pdf_name,
        source_pdf_sha256=pdf_hash,
        source_io_name=io_name,
        source_io_sha256=io_hash,
        processing_version=version,
        draft=draft,
        verification_state=verification_state,
        failed_check=failed_check,
        hidden_count=hidden_count,
        disclaimer=cfg.DISCLAIMER,
        certification_warning=CERTIFICATION_WARNING,
    )


def render_png_and_pdf(svg: str) -> tuple[bytes, bytes]:
    """Turn SVG text into a PNG and a one-page PDF. The SVG file is not the original PDF."""
    document = fitz.open(stream=svg.encode("utf-8"), filetype="svg")
    try:
        page = document[0]
        png = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes("png")
        pdf_doc = fitz.open()
        pdf_page = pdf_doc.new_page(width=page.rect.width, height=page.rect.height + 36)
        pdf_page.insert_image(fitz.Rect(0, 0, page.rect.width, page.rect.height), stream=png)
        pdf_page.insert_text((12, page.rect.height + 18), cfg.DISCLAIMER, fontsize=8)
        pdf_bytes = pdf_doc.tobytes()
        pdf_doc.close()
    finally:
        document.close()
    return png, pdf_bytes


def workbook_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    book = Workbook()
    first = True
    for name, rows in sheets.items():
        sheet = book.active if first else book.create_sheet(name)
        first = False
        sheet.title = name[:31]
        for row_index, row in enumerate(rows, start=1):
            for column, value in enumerate(row, start=1):
                sheet.cell(row_index, column, excel_cell(value))
        for column in range(1, max((len(row) for row in rows), default=1) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 22
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def guard_size(payload: bytes) -> bytes:
    if len(payload) > MAX_EXPORT_BYTES:
        raise AppError(413, "EXPORT_TOO_LARGE", "This export is larger than the limit. No file was returned.", {"bytes": len(payload), "limit": MAX_EXPORT_BYTES})
    return payload


def cache_export(cache_dir: str, key: str, name: str, payload: bytes) -> Path:
    """Store a copy under a short name. A full hash is too long for some Windows paths."""
    short = hashlib.sha256(key.encode("utf-8")).hexdigest()[:32]
    path = Path(cache_dir) / "exports" / short / name
    try:
        if not path.is_file():
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(payload)
    except OSError as exc:
        logger.info("export cache was not written (%s)", type(exc).__name__)
    return path


def source_hash(upload_dir: str, job_id: str) -> str:
    folder = job_folder(upload_dir, job_id)
    return hashlib.sha256((folder / PDF_NAME).read_bytes()).hexdigest()


def graph_json(graph: DiagramGraph) -> bytes:
    return json.dumps(graph.model_dump(mode="json"), indent=2).encode("utf-8")
